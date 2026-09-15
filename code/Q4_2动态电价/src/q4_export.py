from __future__ import annotations
from copy import copy
from openpyxl import load_workbook
from datetime import datetime,timedelta
import numpy as np

PERIODS=['0:00-4:00','4:00-8:00','8:00-12:00','12:00-16:00','16:00-20:00','20:00-24:00']

def _emergency_runs(arr,thr=1e-7):
    arr=np.asarray(arr,float)
    out=[];i=0
    while i<144:
        if arr[i]<=thr: i+=1;continue
        j=i
        while j+1<144 and arr[j+1]>thr: j+=1
        def tm(k):
            minutes=(k+1)*10
            if minutes==1440:return '0:00+1'
            return f'{minutes//60}:{minutes%60:02d}'
        start='0:00' if i==0 else tm(i-1)
        end=tm(j)
        out.append((f'{start}-{end}',float(arr[i:j+1].sum())))
        i=j+1
    return out

def export_result42(template,out,results,actual_prices):
    wb=load_workbook(template)
    ws=wb['计划购电量']
    for i,r in enumerate(results,2):
        ws.cell(i,1).value=datetime.fromisoformat(r['date'])
        for t,v in enumerate(r['grid'],2): ws.cell(i,t).value=float(v)
        ws.cell(i,146).value=float(np.sum(r['grid']))
        ws.cell(i,147).value=float(np.dot(actual_prices[r['day_idx']],r['grid']))

    ws=wb['充放电量']
    # remove template example rows, preserve header/style row
    if ws.max_row>1: ws.delete_rows(2,ws.max_row-1)
    row=2
    for r in results:
        date=datetime.fromisoformat(r['date'])
        for k,period in enumerate(PERIODS):
            a=k*24;b=(k+1)*24
            ws.cell(row,1).value=date if k==0 else None
            ws.cell(row,2).value=period
            ws.cell(row,3).value=float(np.sum(r['charge'][a:b]))
            ws.cell(row,4).value=float(np.sum(r['discharge'][a:b]))
            if k==0:
                ws.cell(row,5).value='0:00';ws.cell(row,6).value=float(r['soc_start'])
            elif k==1:
                ws.cell(row,5).value='24:00';ws.cell(row,6).value=float(r['soc_end'])
            row+=1

    ws=wb['紧急购电量']
    if ws.max_row>1: ws.delete_rows(2,ws.max_row-1)
    row=2
    for r in results:
        runs=_emergency_runs(r['emergency'])
        if not runs:
            continue
        for k,(period,qty) in enumerate(runs):
            ws.cell(row,1).value=datetime.fromisoformat(r['date']) if k==0 else None
            ws.cell(row,2).value=period;ws.cell(row,3).value=qty;row+=1
    wb.save(out)

def validate_result42(path,results,prices):
    wb=load_workbook(path,read_only=True,data_only=True)
    ws=wb['计划购电量']
    if ws.max_row!=335 or ws.max_column!=147:
        raise ValueError(f'计划购电量尺寸异常 {ws.max_row}x{ws.max_column}')
    for i,(row,r) in enumerate(zip(ws.iter_rows(min_row=2,max_row=335,values_only=True),results),2):
        vals=np.asarray(row[1:145],float)
        if abs(vals.sum()-float(row[145]))>1e-5: raise ValueError(f'第{i}行总量错误')
        want=float(np.dot(prices[r["day_idx"]],vals))
        if abs(want-float(row[146]))>1e-3: raise ValueError(f'第{i}行费用错误')
    ws=wb['充放电量']
    if ws.max_row!=1+334*6: raise ValueError(f'充放电量行数错误 {ws.max_row}')
    wb.close();return True
