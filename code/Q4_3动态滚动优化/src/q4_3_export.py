from __future__ import annotations
import shutil
import numpy as np
from openpyxl import load_workbook

PERIODS=['0:00-4:00','4:00-8:00','8:00-12:00','12:00-16:00','16:00-20:00','20:00-24:00']

def _emergency_blocks(vals, headers, eps=1e-7):
    out=[]; i=0
    while i<len(vals):
        if vals[i] <= eps: i+=1; continue
        j=i; total=0.0
        while j<len(vals) and vals[j]>eps:
            total += float(vals[j]); j+=1
        out.append((f"{headers[i].split('-')[0]}-{headers[j-1].split('-')[-1]}", total))
        i=j
    return out


def export_result43(template_path,out_path,results,actual_prices):
    shutil.copy2(template_path,out_path)
    wb=load_workbook(out_path)
    wp=wb['计划购电量']; wa=wb['调整购电量']
    wbatt=wb['充放电量']; we=wb['紧急购电量']
    headers=[wp.cell(1,c).value for c in range(2,146)]

    for i,r in enumerate(results,start=2):
        p=actual_prices[r['day_idx']]
        wp.cell(i,1,r['date']); wa.cell(i,1,r['date'])
        for t in range(144):
            wp.cell(i,2+t,float(r['original_plan'][t]))
            wa.cell(i,2+t,float(r['final_plan'][t]))
        wp.cell(i,146,float(np.sum(r['original_plan'])))
        wp.cell(i,147,float(np.dot(p,r['original_plan'])))
        wa.cell(i,146,float(np.sum(r['final_plan'])))
        wa.cell(i,147,float(r['cost']['total_cost']))

    if wbatt.max_row>1: wbatt.delete_rows(2,wbatt.max_row-1)
    rr=2
    for r in results:
        for b,period in enumerate(PERIODS):
            a=b*24;z=(b+1)*24
            wbatt.cell(rr,1,r['date'] if b==0 else None)
            wbatt.cell(rr,2,period)
            wbatt.cell(rr,3,float(np.sum(r['charge'][a:z])))
            wbatt.cell(rr,4,float(np.sum(r['discharge'][a:z])))
            if b==0:
                wbatt.cell(rr,5,'0:00');wbatt.cell(rr,6,float(r['soc_start']))
            elif b==1:
                wbatt.cell(rr,5,'24:00');wbatt.cell(rr,6,float(r['soc_end']))
            rr+=1

    if we.max_row>1: we.delete_rows(2,we.max_row-1)
    rr=2
    for r in results:
        blocks=_emergency_blocks(r['emergency'],headers)
        for k,(period,val) in enumerate(blocks):
            we.cell(rr,1,r['date'] if k==0 else None)
            we.cell(rr,2,period);we.cell(rr,3,float(val));rr+=1
    if rr==2:
        we.cell(2,1,results[0]['date']);we.cell(2,2,'无');we.cell(2,3,0.0)
    wb.save(out_path)

def validate_result43_template(template_path):
    wb=load_workbook(template_path,read_only=True,data_only=False)
    required=['计划购电量','调整购电量','充放电量','紧急购电量']
    if wb.sheetnames[:4] != required:
        raise ValueError(f'模板工作表异常: {wb.sheetnames}')
    for name in ('计划购电量','调整购电量'):
        ws=wb[name]
        if ws.max_row<335 or ws.max_column!=147:
            raise ValueError(f'{name}尺寸异常 {ws.max_row}x{ws.max_column}')
    wb.close();return True

def validate_result43_output(out_path,results,actual_prices,tol=1e-5):
    wb=load_workbook(out_path,read_only=True,data_only=True)
    if len(results)!=334: raise ValueError('全年结果不是334天')
    wp=wb['计划购电量'];wa=wb['调整购电量']
    if wp.max_row!=335 or wp.max_column!=147: raise ValueError('计划购电量尺寸异常')
    if wa.max_row!=335 or wa.max_column!=147: raise ValueError('调整购电量尺寸异常')
    plan_rows=list(wp.iter_rows(min_row=2,max_row=335,values_only=True))
    adj_rows=list(wa.iter_rows(min_row=2,max_row=335,values_only=True))
    for idx,(pr,ar,r) in enumerate(zip(plan_rows,adj_rows,results),start=2):
        p=actual_prices[r['day_idx']]
        pv=np.asarray([(x or 0.0) for x in pr[1:145]],float)
        av=np.asarray([(x or 0.0) for x in ar[1:145]],float)
        if abs(float(pr[145])-pv.sum())>tol: raise ValueError(f'计划第{idx}行总量错')
        if abs(float(ar[145])-av.sum())>tol: raise ValueError(f'调整第{idx}行总量错')
        if abs(float(pr[146])-float(np.dot(p,pv)))>1e-3: raise ValueError(f'计划第{idx}行费用错')
        if abs(float(ar[146])-float(r['cost']['total_cost']))>1e-3: raise ValueError(f'调整第{idx}行费用错')
    ws=wb['充放电量']
    if ws.max_row != 1+334*6 or ws.max_column!=6:
        raise ValueError(f'充放电量尺寸异常 {ws.max_row}x{ws.max_column}')
    wb.close();return True
